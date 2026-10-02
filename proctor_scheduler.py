# -*- coding: utf-8 -*-
"""
감독관 배정 시스템
- 입력: 시험실 수, 감독관 수
- 규칙: 1~3교시 = 시험실당 정 1 + 부 1 / 4교시 = 시험실당 정 1 + 부 2
        감독관 전원 4교시 참여, 1인당 총 3회 (감독관 수 = 시험실 수 x 3 일 때 정확히 성립)
        정감독 고른 분포, 연속 3회는 불가피한 인원만 + 그 인원의 정감독 최소화
- 알고리즘: 최대유량(Max-Flow)으로 '참여 교시' 와 '정감독 배정'을 각각 정확히 풀고,
            시험실 배치는 랜덤 탐색으로 같은 짝/같은 방 반복을 최소화
"""
import csv
import random
import tkinter as tk
import tkinter.font as tkfont
from collections import Counter, deque
from itertools import combinations
from tkinter import filedialog, messagebox, ttk

PERIODS = 4  # 기본 교시 수
MAX_PERIODS = 5  # 5교시(선택) 포함

# ───────────────────────────── 배정 로직 ─────────────────────────────


class MaxFlow:
    """간단한 Edmonds-Karp 최대유량 (간선 id 방식 → 인접 리스트 셔플로 무작위성 확보)"""

    def __init__(self, n):
        self.adj = [[] for _ in range(n)]
        self.to = []
        self.cap = []

    def add(self, u, v, c):
        self.adj[u].append(len(self.to))
        self.to.append(v)
        self.cap.append(c)
        self.adj[v].append(len(self.to))
        self.to.append(u)
        self.cap.append(0)
        return len(self.to) - 2

    def run(self, s, t, rng):
        for a in self.adj:
            rng.shuffle(a)
        flow = 0
        while True:
            prev = [-1] * len(self.adj)
            prev[s] = -2
            q = deque([s])
            while q and prev[t] == -1:
                u = q.popleft()
                for e in self.adj[u]:
                    v = self.to[e]
                    if self.cap[e] > 0 and prev[v] == -1:
                        prev[v] = e
                        q.append(v)
            if prev[t] == -1:
                break
            v = t
            while v != s:
                e = prev[v]
                self.cap[e] -= 1
                self.cap[e ^ 1] += 1
                v = self.to[e ^ 1]
            flow += 1
        return flow


def is_consec3(periods):
    s = set(periods)
    return any({a, a + 1, a + 2} <= s for a in range(MAX_PERIODS - 2))


def build_attendance(R, N, rng):
    """교시별 참여자 결정. 반환: attend[p] = set(감독관 index)"""
    demand = [2 * R, 2 * R, 2 * R, 3 * R]
    total = sum(demand)
    base, rem = divmod(total, N)
    caps = [base] * N
    for i in rng.sample(range(N), rem):
        caps[i] += 1

    S, T = 0, 1
    mf = MaxFlow(2 + PERIODS + N)
    edges = {}
    for p in range(PERIODS):
        mf.add(S, 2 + p, demand[p])
        for i in range(N):
            edges[(p, i)] = mf.add(2 + p, 2 + PERIODS + i, 1)
    for i in range(N):
        mf.add(2 + PERIODS + i, T, caps[i])
    if mf.run(S, T, rng) < total:
        return None
    attend = [set() for _ in range(PERIODS)]
    for (p, i), e in edges.items():
        if mf.cap[e] == 0:
            attend[p].add(i)
    return attend


def count_consec(attend, N):
    n = 0
    for i in range(N):
        if is_consec3([p for p in range(PERIODS) if i in attend[p]]):
            n += 1
    return n


def assign_roles(R, N, attend, rng, reserve=frozenset()):
    """정감독 결정. 연속3회 인원은 최소 횟수, 나머지는 균등하게."""
    per_person = [[p for p in range(PERIODS) if i in attend[p]] for i in range(N)]
    cons = [is_consec3(x) for x in per_person]
    total = R * PERIODS
    base = total // N

    for attempt in range(60):
        target = [min(base, len(per_person[i])) for i in range(N)]
        pool_a = [i for i in range(N) if not cons[i] and i not in reserve]
        pool_b = [i for i in range(N) if not cons[i] and i in reserve]
        pool_c = [i for i in range(N) if cons[i]]
        rng.shuffle(pool_a)
        rng.shuffle(pool_b)
        rng.shuffle(pool_c)
        pool = pool_a + pool_b  # 5교시 후보는 추가 정감독을 맨 나중에 받음
        need = total - sum(target)
        guard = 0
        while need > 0 and guard < 50:
            for group in (pool, pool_c):
                for i in group:
                    if need > 0 and target[i] < len(per_person[i]):
                        target[i] += 1
                        need -= 1
                if need <= 0:
                    break
            guard += 1
        slack = 1 if attempt >= 30 else 0  # 막히면 약간 완화

        S, T = 0, 1
        mf = MaxFlow(2 + PERIODS + N)
        edges = {}
        for p in range(PERIODS):
            mf.add(S, 2 + p, R)
            for i in attend[p]:
                edges[(p, i)] = mf.add(2 + p, 2 + PERIODS + i, 1)
        for i in range(N):
            mf.add(2 + PERIODS + i, T, target[i] + (slack if not cons[i] else 0))
        if mf.run(S, T, rng) == total:
            jung = [set() for _ in range(PERIODS)]
            for (p, i), e in edges.items():
                if mf.cap[e] == 0:
                    jung[p].add(i)
            return jung
    return None


def assign_rooms(R, attend, jung, rng, tries=400):
    """시험실 배치. 같은 짝(정-부, 부-부) 반복 및 같은 방 반복을 최소화."""
    best, best_cost = None, None
    for _ in range(tries):
        layout = []
        for p in range(PERIODS):
            j = list(jung[p])
            rng.shuffle(j)
            b = [i for i in attend[p] if i not in jung[p]]
            rng.shuffle(b)
            k = len(b) // R
            layout.append([(j[r], b[r * k:(r + 1) * k]) for r in range(R)])

        pair_cnt, room_cnt = Counter(), Counter()
        for p in range(PERIODS):
            for r, (j, bs) in enumerate(layout[p]):
                members = [j] + bs
                for a, c in combinations(sorted(members), 2):
                    pair_cnt[(a, c)] += 1
                for m in members:
                    room_cnt[(m, r)] += 1
        cost = sum((v - 1) * 3 for v in pair_cnt.values() if v > 1) + \
            sum(v - 1 for v in room_cnt.values() if v > 1)
        if best_cost is None or cost < best_cost:
            best, best_cost = layout, cost
            if cost == 0:
                break
    return best


def assign_period5(R5, N, layout, elig, rng, tries=400):
    """5교시: 3·4교시 연속 참여자를 제외한 인원 중에서 시험실당 정 1 + 부 1.
    정감독은 1~4교시 정감독 횟수가 적은 사람(= 다른 교시 정감독과 최대한 안 겹치게) 우선."""
    jcnt, tot = [0] * N, [0] * N
    for p in range(PERIODS):
        for j, bs in layout[p]:
            jcnt[j] += 1
            tot[j] += 1
            for b in bs:
                tot[b] += 1
    p4_jung = {j for j, _ in layout[3]}
    ranked = sorted(elig, key=lambda i: (jcnt[i], tot[i], i in p4_jung, rng.random()))
    jung5 = ranked[:R5]
    rest = sorted(ranked[R5:], key=lambda i: (tot[i], rng.random()))
    bu5 = rest[:R5]

    prev_pairs, prev_room = Counter(), set()
    for p in range(PERIODS):
        for r, (j, bs) in enumerate(layout[p]):
            mem = [j] + bs
            for a, c in combinations(sorted(mem), 2):
                prev_pairs[(a, c)] += 1
            for m in mem:
                prev_room.add((m, r))

    best, best_cost = None, None
    for _ in range(tries):
        j5, b5 = jung5[:], bu5[:]
        rng.shuffle(j5)
        rng.shuffle(b5)
        cost = 0
        for r in range(R5):
            a, c = sorted((j5[r], b5[r]))
            cost += prev_pairs.get((a, c), 0) * 3
            cost += ((j5[r], r) in prev_room) + ((b5[r], r) in prev_room)
        if best_cost is None or cost < best_cost:
            best, best_cost = [(j5[r], [b5[r]]) for r in range(R5)], cost
            if cost == 0:
                break
    return best


def note_of(s):
    notes = []
    if s["consec3"]:
        notes.append("연속 3회")
    if s["p5jung"] and s["jung"] >= 2:
        notes.append("5교시 정 겸임")
    return " · ".join(notes)


def solve(R, N, R5=0):
    rng = random.Random()
    best_att, best_key = None, None
    for _ in range(40):
        att = build_attendance(R, N, rng)
        if att is None:
            continue
        c = count_consec(att, N)
        elig_n = N - len(att[2] & att[3])
        key = (elig_n < 2 * R5, c, -elig_n)
        if best_key is None or key < best_key:
            best_att, best_key = att, key
    if best_att is None:
        return None

    elig = {i for i in range(N) if not (i in best_att[2] and i in best_att[3])}
    if R5 > 0 and len(elig) < 2 * R5:
        raise ValueError(
            f"5교시에는 3·4교시 연속 참여자를 제외한 {len(elig)}명만 투입할 수 있어\n"
            f"5교시 시험실은 최대 {len(elig) // 2}개까지 가능합니다.")

    jung = None
    if R5 > 0:
        jung = assign_roles(R, N, best_att, rng, reserve=elig)
    if jung is None:
        jung = assign_roles(R, N, best_att, rng)
    if jung is None:
        return None
    layout = assign_rooms(R, best_att, jung, rng)
    if R5 > 0:
        layout.append(assign_period5(R5, N, layout, elig, rng))

    stats = []
    for i in range(N):
        stats.append({"jung": 0, "bu": 0, "cells": ["-"] * MAX_PERIODS,
                      "periods": [], "p5jung": False})
    for p in range(len(layout)):
        for r, (j, bs) in enumerate(layout[p]):
            stats[j]["jung"] += 1
            stats[j]["p5jung"] = stats[j]["p5jung"] or p == 4
            stats[j]["cells"][p] = f"{r + 1}실 · 정"
            stats[j]["periods"].append(p)
            for b in bs:
                stats[b]["bu"] += 1
                stats[b]["cells"][p] = f"{r + 1}실 · 부"
                stats[b]["periods"].append(p)
    for s in stats:
        s["consec3"] = is_consec3(s["periods"])
    overlap = sum(1 for s in stats if s["p5jung"] and s["jung"] >= 2)
    return {"R": R, "N": N, "R5": R5, "layout": layout, "stats": stats,
            "attend": best_att, "p5_overlap": overlap}


# ───────────────────────────── 엑셀 출력 ─────────────────────────────


def write_excel(path, res):
    """시간표 스타일의 컬러 엑셀(.xlsx) 저장. (openpyxl 필요)"""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as CL
    from openpyxl.worksheet.properties import PageSetupProperties

    R, N, R5 = res["R"], res["N"], res["R5"]
    layout, stats = res["layout"], res["stats"]
    npd = 5 if R5 else 4
    FONT, DARK = "Malgun Gothic", "1E1B4B"
    PAL = ["4F46E5", "0D9488", "D97706", "E11D48", "7C3AED"]  # 교시별 대표색

    def tint(h, t):
        rgb = [int(h[i:i + 2], 16) for i in (0, 2, 4)]
        return "%02X%02X%02X" % tuple(round(c + (255 - c) * t) for c in rgb)

    def fill(h):
        return PatternFill("solid", start_color=h, end_color=h)

    def font(**kw):
        return Font(name=FONT, **kw)

    thin = Side(style="thin", color="D1D5DB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def banner(ws, title, sub, ncols):
        for r, (txt, size, bold, color, h) in enumerate(
                [(title, 20, True, "FFFFFF", 40), (sub, 10, False, "C7D2FE", 24)], start=1):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
            c = ws.cell(r, 1, txt)
            c.font = font(size=size, bold=bold, color=color)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            for col in range(1, ncols + 1):
                ws.cell(r, col).fill = fill(DARK)
            ws.row_dimensions[r].height = h
        ws.row_dimensions[3].height = 10
        ws.sheet_view.showGridLines = False

    def page(ws):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    sub = (f"시험실 {R}개 · 감독관 {N}명" + (f" · 5교시 {R5}개 시험실" if R5 else "")
           + "        ■ 진한 색 = 정감독    ■ 연한 색 = 부감독")

    # ───────── 시트 1: 배정 시간표 ─────────
    wb = Workbook()
    ws = wb.active
    ws.title = "배정 시간표"
    ws.sheet_properties.tabColor = PAL[0]
    ncols = 1 + sum(3 if p == 3 else 2 for p in range(npd))
    banner(ws, "서창고 수능감독관 배정 시간표", sub, ncols)

    ws.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
    for r in (4, 5):
        ws.cell(r, 1).fill = fill(DARK)
        ws.cell(r, 1).border = box
    ws.cell(4, 1, "시험실").font = font(size=11, bold=True, color="FFFFFF")
    ws.cell(4, 1).alignment = center

    col = 2
    col_of = {}
    for p in range(npd):
        w = 3 if p == 3 else 2
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + w - 1)
        for k in range(w):
            c4 = ws.cell(4, col + k)
            c4.fill, c4.border = fill(PAL[p]), box
            c5 = ws.cell(5, col + k, "정" if k == 0 else "부")
            c5.fill, c5.border = fill(tint(PAL[p], 0.55 if k == 0 else 0.85)), box
            c5.font = font(size=10, bold=True, color=DARK)
            c5.alignment = center
            ws.column_dimensions[CL(col + k)].width = 9
        ws.cell(4, col, f"{p + 1}교시").font = font(size=12, bold=True, color="FFFFFF")
        ws.cell(4, col).alignment = center
        col_of[p] = col
        col += w
    ws.row_dimensions[4].height = 26
    ws.row_dimensions[5].height = 20
    ws.column_dimensions["A"].width = 12

    rows = max(R, R5)
    for r in range(rows):
        row = 6 + r
        c = ws.cell(row, 1, f"{r + 1}실")
        c.font, c.fill, c.border, c.alignment = font(size=11, bold=True, color=DARK), fill("EEF2FF"), box, center
        for p in range(npd):
            w = 3 if p == 3 else 2
            ppl = None
            if p < len(layout) and r < len(layout[p]):
                j, bs = layout[p][r]
                ppl = [j] + bs
            for k in range(w):
                cell = ws.cell(row, col_of[p] + k)
                cell.border, cell.alignment = box, center
                if ppl is None:
                    cell.value = "-"
                    cell.font, cell.fill = font(size=10, color="9CA3AF"), fill("F3F4F6")
                else:
                    cell.value = ppl[k] + 1
                    cell.font = font(size=12, bold=(k == 0), color=DARK)
                    cell.fill = fill(tint(PAL[p], 0.55 if k == 0 else 0.88))
        ws.row_dimensions[row].height = 26
    ws.freeze_panes = "B6"
    page(ws)

    # ───────── 시트 2: 감독관별 통계 ─────────
    ws2 = wb.create_sheet("감독관별 통계")
    ws2.sheet_properties.tabColor = PAL[1]
    heads = ["감독관", "정감독", "부감독", "합계"] + [f"{p + 1}교시" for p in range(npd)] + ["비고"]
    n2 = len(heads)
    banner(ws2, "감독관별 배정 통계",
           "교시별 배정 시험실 표시   ■ 진한 색 = 정감독   ■ 연한 색 = 부감독   ■ 노란 이름 = 연속 3회 불가피 인원", n2)
    for ci, h in enumerate(heads, start=1):
        c = ws2.cell(4, ci, h)
        base = PAL[ci - 5] if 5 <= ci < 5 + npd else DARK
        c.fill, c.border, c.alignment = fill(base), box, center
        c.font = font(size=11, bold=True, color="FFFFFF")
    ws2.row_dimensions[4].height = 26
    widths = [11, 10, 10, 9] + [13] * npd + [22]
    for ci, wd in enumerate(widths, start=1):
        ws2.column_dimensions[CL(ci)].width = wd

    for i, st in enumerate(stats):
        row = 5 + i
        vals = [f"{i + 1}번", st["jung"], st["bu"], st["jung"] + st["bu"]]
        for ci, v in enumerate(vals, start=1):
            c = ws2.cell(row, ci, v)
            c.border, c.alignment = box, center
            c.font = font(size=11, bold=(ci in (1, 2)), color=DARK)
            c.fill = fill("FEF3C7" if (ci == 1 and st["consec3"]) else ("EEF2FF" if ci == 1 else "FFFFFF"))
        for p in range(npd):
            txt = st["cells"][p]
            c = ws2.cell(row, 5 + p, txt)
            c.border, c.alignment = box, center
            if txt.endswith("정"):
                c.fill, c.font = fill(tint(PAL[p], 0.55)), font(size=11, bold=True, color=DARK)
            elif txt.endswith("부"):
                c.fill, c.font = fill(tint(PAL[p], 0.88)), font(size=11, color=DARK)
            else:
                c.value = "-"
                c.fill, c.font = fill("F3F4F6"), font(size=10, color="9CA3AF")
        c = ws2.cell(row, n2, note_of(st))
        c.border, c.alignment = box, center
        c.font = font(size=10, bold=True, color="B45309")
        if st["consec3"]:
            c.fill = fill("FEF3C7")
        ws2.row_dimensions[row].height = 22
    last = 4 + N
    ws2.conditional_formatting.add(
        f"B5:B{last}", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                      end_type="max", end_color="A5B4FC"))
    ws2.auto_filter.ref = f"A4:{CL(n2)}{last}"
    ws2.freeze_panes = "B5"
    page(ws2)

    # ───────── 시트 3: 요약 ─────────
    ws3 = wb.create_sheet("요약")
    ws3.sheet_properties.tabColor = PAL[2]
    banner(ws3, "배정 요약", sub, 14)
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 14
    ws3.column_dimensions["C"].width = 4
    ws3.column_dimensions["D"].width = 18
    ws3.column_dimensions["E"].width = 14

    def put(r, c, v, bold=False, bg=None, color=DARK, al=center):
        cell = ws3.cell(r, c, v)
        cell.font, cell.border, cell.alignment = font(size=11, bold=bold, color=color), box, al
        if bg:
            cell.fill = fill(bg)

    cons_n = sum(1 for x in stats if x["consec3"])
    metrics = [("시험실 수", f"{R}개"), ("감독관 수", f"{N}명"),
               ("5교시 시험실 수", f"{R5}개" if R5 else "미실시"),
               ("총 배정 슬롯", f"{sum(x['jung'] + x['bu'] for x in stats)}석"),
               ("4교시 참여 인원", f"{len(res['attend'][3])} / {N}명"),
               ("연속 3회 인원", f"{cons_n}명"),
               ("5교시 정감독 중 타 교시 정 겸임", f"{res['p5_overlap']} / {R5}명" if R5 else "-")]
    put(4, 1, "항목", True, DARK, "FFFFFF")
    put(4, 2, "값", True, DARK, "FFFFFF")
    for k, (a, b) in enumerate(metrics):
        put(5 + k, 1, a, True, "EEF2FF", al=Alignment(horizontal="left", vertical="center", indent=1))
        put(5 + k, 2, b, False)
    ws3.column_dimensions["A"].width = 34

    def dist(r0, title, key):
        cnt = Counter(key(x) for x in stats)
        put(r0, 4, title, True, DARK, "FFFFFF")
        put(r0, 5, "인원", True, DARK, "FFFFFF")
        for k, v in enumerate(sorted(cnt)):
            put(r0 + 1 + k, 4, f"{v}회", True, "EEF2FF")
            put(r0 + 1 + k, 5, cnt[v])
        return r0, r0 + len(cnt)

    j0, j1 = dist(4, "정감독 횟수", lambda x: x["jung"])
    t0, t1 = dist(max(j1 + 2, 9), "총 배정 횟수", lambda x: x["jung"] + x["bu"])
    ch = BarChart()
    ch.type, ch.title, ch.style = "col", "정감독 횟수별 인원", 10
    ch.add_data(Reference(ws3, min_col=5, min_row=j0, max_row=j1), titles_from_data=True)
    ch.set_categories(Reference(ws3, min_col=4, min_row=j0 + 1, max_row=j1))
    ch.legend, ch.height, ch.width = None, 7.5, 11
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorGridlines = None
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.series[0].graphicalProperties.solidFill = PAL[0]
    ws3.add_chart(ch, "G4")
    page(ws3)

    wb.save(path)


# ───────────────────────────── GUI ─────────────────────────────

NAVY = "#1E1B4B"
ACCENT = "#4F46E5"
ACCENT_DARK = "#4338CA"
ACCENT_LIGHT = "#E0E7FF"
BG = "#F3F4F8"
CARD = "#FFFFFF"
BORDER = "#E5E7EB"
TEXT = "#111827"
MUTED = "#6B7280"
WARN_BG = "#FEF3C7"
STRIPE = "#F8FAFC"


def pick_font(root):
    fams = set(tkfont.families(root))
    for f in ("Pretendard", "Malgun Gothic", "Apple SD Gothic Neo",
              "NanumGothic", "Noto Sans CJK KR", "Segoe UI"):
        if f in fams:
            return f
    return "TkDefaultFont"


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("서창고 수능감독관 배정 시스템")
        self.geometry("1200x780")
        self.minsize(1000, 650)
        self.configure(bg=BG)
        self.F = pick_font(self)
        self.result = None
        self.rooms_var = tk.StringVar(value="14")
        self.staff_var = tk.StringVar(value="42")
        self.rooms5_var = tk.StringVar(value="2")
        self._setup_style()
        self._build()
        self.rooms_var.trace_add("write", lambda *_: self._update_hint())
        self.staff_var.trace_add("write", lambda *_: self._update_hint())
        self.rooms5_var.trace_add("write", lambda *_: self._update_hint())
        self._update_hint()
        self.generate()

    # ── 스타일 ──
    def _setup_style(self):
        F = self.F
        st = ttk.Style(self)
        st.theme_use("clam")
        st.configure("Treeview", background="white", fieldbackground="white",
                     foreground=TEXT, rowheight=32, font=(F, 10), borderwidth=0)
        st.configure("Treeview.Heading", background=NAVY, foreground="white",
                     font=(F, 10, "bold"), relief="flat", padding=(8, 9))
        st.map("Treeview.Heading", background=[("active", ACCENT)])
        st.map("Treeview", background=[("selected", ACCENT_LIGHT)],
               foreground=[("selected", TEXT)])
        st.configure("TNotebook", background=BG, borderwidth=0)
        st.configure("TNotebook.Tab", font=(F, 10, "bold"),
                     background="#E5E7EB", foreground="#4B5563", borderwidth=0)
        st.configure("Accent.TButton", background=ACCENT, foreground="white",
                     font=(F, 10, "bold"), padding=(20, 9), borderwidth=0,
                     focusthickness=0, relief="flat")
        st.map("Accent.TButton", background=[("active", ACCENT_DARK)])
        st.configure("Ghost.TButton", background="white", foreground=ACCENT,
                     font=(F, 10, "bold"), padding=(18, 8), bordercolor=ACCENT,
                     lightcolor="white", darkcolor="white", focusthickness=0)
        st.map("Ghost.TButton", background=[("active", ACCENT_LIGHT)])
        st.configure("Green.TButton", background="#059669", foreground="white",
                     font=(F, 10, "bold"), padding=(20, 9), borderwidth=0,
                     focusthickness=0, relief="flat")
        st.map("Green.TButton", background=[("active", "#047857")])
        st.configure("TSpinbox", padding=6, arrowsize=14, font=(F, 11))
        st.configure("TEntry", padding=6)

    # ── 레이아웃 ──
    def _build(self):
        F = self.F
        # 헤더
        header = tk.Frame(self, bg=NAVY)
        header.pack(fill="x")
        tk.Label(header, text="서창고 수능감독관 배정 시스템", bg=NAVY, fg="white",
                 font=(F, 20, "bold")).pack(anchor="w", padx=28, pady=(18, 0))
        tk.Label(header, text="시험실 수와 감독관 수만 입력하면 교시별 배정표와 감독관별 통계를 자동으로 만들어 드립니다",
                 bg=NAVY, fg="#C7D2FE", font=(F, 10)).pack(anchor="w", padx=28, pady=(2, 18))

        body = tk.Frame(self, bg=BG)
        body.pack(fill="both", expand=True, padx=24, pady=18)

        # 입력 카드
        card = self._card(body)
        card.pack(fill="x")
        row = tk.Frame(card, bg=CARD)
        row.pack(fill="x", padx=20, pady=16)

        def field(parent, label, var, hi):
            f = tk.Frame(parent, bg=CARD)
            tk.Label(f, text=label, bg=CARD, fg=MUTED, font=(F, 9, "bold")).pack(anchor="w")
            ttk.Spinbox(f, from_=1, to=hi, textvariable=var, width=8,
                        font=(F, 12)).pack(anchor="w", pady=(4, 0))
            return f

        field(row, "시험실 수", self.rooms_var, 200).pack(side="left", padx=(0, 22))
        field(row, "감독관 수", self.staff_var, 1000).pack(side="left", padx=(0, 22))
        field(row, "5교시 시험실 수 (0 = 없음)", self.rooms5_var, 200).pack(side="left", padx=(0, 26))
        ttk.Button(row, text="배정 생성", style="Accent.TButton",
                   command=self.generate).pack(side="left", padx=(0, 10), pady=(18, 0))
        ttk.Button(row, text="엑셀 저장 (시간표)", style="Green.TButton",
                   command=self.export_excel).pack(side="left", padx=(0, 10), pady=(18, 0))
        ttk.Button(row, text="CSV 저장", style="Ghost.TButton",
                   command=self.export_csv).pack(side="left", pady=(18, 0))
        self.hint = tk.Label(card, text="", bg=CARD, fg=MUTED, font=(F, 9),
                             justify="left", anchor="w")
        self.hint.pack(anchor="w", padx=20, pady=(0, 12))

        # 요약 타일
        tiles = tk.Frame(body, bg=BG)
        tiles.pack(fill="x", pady=14)
        self.tile_vals = {}
        for k, title in (("slots", "총 배정 슬롯"), ("jung", "정감독 횟수 분포"),
                         ("cons", "연속 3회 인원"), ("p4", "4교시 참여"),
                         ("p5", "5교시 정감독 중 타 교시 정 겸임")):
            t = self._card(tiles)
            t.pack(side="left", fill="x", expand=True, padx=(0, 12))
            tk.Label(t, text=title, bg=CARD, fg=MUTED, font=(F, 9, "bold")).pack(
                anchor="w", padx=16, pady=(12, 0))
            v = tk.Label(t, text="-", bg=CARD, fg=ACCENT, font=(F, 18, "bold"))
            v.pack(anchor="w", padx=16, pady=(0, 12))
            self.tile_vals[k] = v

        # 탭
        nb = ttk.Notebook(body)
        nb.pack(fill="both", expand=True)
        tab1 = tk.Frame(nb, bg="white")
        tab2 = tk.Frame(nb, bg="white")
        nb.add(tab1, text="  교시별 배정표  ")
        nb.add(tab2, text="  감독관별 통계  ")

        cols1 = ("room", "p1", "p2", "p3", "p4", "p5")
        self.cols1 = cols1
        self.tree1 = self._tree(tab1, cols1)
        heads = {"room": ("시험실", 80), "p1": ("1교시  (정 · 부)", 170),
                 "p2": ("2교시  (정 · 부)", 170), "p3": ("3교시  (정 · 부)", 170),
                 "p4": ("4교시  (정 · 부 · 부)", 230), "p5": ("5교시  (정 · 부)", 170)}
        for c in cols1:
            self.tree1.heading(c, text=heads[c][0])
            self.tree1.column(c, width=heads[c][1], anchor="center")

        # 통계 탭 (검색 바 + 표)
        bar = tk.Frame(tab2, bg="white")
        bar.pack(fill="x", padx=12, pady=(10, 4))
        tk.Label(bar, text="감독관 번호 찾기", bg="white", fg=MUTED,
                 font=(F, 9, "bold")).pack(side="left")
        self.find_var = tk.StringVar()
        e = ttk.Entry(bar, textvariable=self.find_var, width=8)
        e.pack(side="left", padx=8)
        e.bind("<Return>", lambda _e: self.find_staff())
        ttk.Button(bar, text="이동", style="Ghost.TButton",
                   command=self.find_staff).pack(side="left")
        tk.Label(bar, text="■ 노란 행 = 연속 3회 불가피 인원", bg="white", fg="#B45309",
                 font=(F, 9)).pack(side="right")

        holder = tk.Frame(tab2, bg="white")
        holder.pack(fill="both", expand=True)
        cols2 = ("no", "jung", "bu", "total", "p1", "p2", "p3", "p4", "p5", "note")
        self.cols2 = cols2
        self.tree2 = self._tree(holder, cols2)
        h2 = {"no": ("감독관", 80), "jung": ("정감독", 70), "bu": ("부감독", 70),
              "total": ("합계", 60), "p1": ("1교시", 90), "p2": ("2교시", 90),
              "p3": ("3교시", 90), "p4": ("4교시", 90), "p5": ("5교시", 90),
              "note": ("비고", 170)}
        for c in cols2:
            self.tree2.heading(c, text=h2[c][0])
            self.tree2.column(c, width=h2[c][1], anchor="center")

        for t in (self.tree1, self.tree2):
            t.tag_configure("odd", background=STRIPE)
            t.tag_configure("even", background="white")
        self.tree2.tag_configure("warn", background=WARN_BG)

        # 상태바
        self.status = tk.Label(self, text="", bg=BG, fg=MUTED, font=(F, 9), anchor="w")
        self.status.pack(fill="x", padx=26, pady=(0, 8))

    def _card(self, parent):
        return tk.Frame(parent, bg=CARD, highlightbackground=BORDER, highlightthickness=1)

    def _tree(self, parent, cols):
        wrap = tk.Frame(parent, bg="white")
        wrap.pack(fill="both", expand=True)
        tree = ttk.Treeview(wrap, columns=cols, show="headings", selectmode="browse")
        ys = ttk.Scrollbar(wrap, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=ys.set)
        tree.pack(side="left", fill="both", expand=True)
        ys.pack(side="right", fill="y")
        return tree

    # ── 동작 ──
    def _update_hint(self):
        try:
            r = int(self.rooms_var.get())
            self.hint.config(text=f"권장 감독관 수 = 시험실 × 3 = {3 * r}명 (전원 4교시 · 1인 3회 성립)   |   "
                                  f"5교시 시험실은 이 경우 최대 {r // 2}개까지 가능 (3·4교시 연속자 제외)")
        except ValueError:
            self.hint.config(text="")

    def generate(self):
        try:
            R = int(self.rooms_var.get())
            N = int(self.staff_var.get())
            R5 = int(self.rooms5_var.get())
        except ValueError:
            messagebox.showwarning("입력 오류", "시험실 수, 감독관 수, 5교시 시험실 수는 숫자로 입력해 주세요.")
            return
        if R < 1 or N < 1 or R5 < 0:
            messagebox.showwarning("입력 오류", "1 이상의 숫자를 입력해 주세요.")
            return
        if N < 3 * R:
            messagebox.showerror(
                "배정 불가",
                f"4교시에는 시험실당 3명씩 총 {3 * R}명이 동시에 필요합니다.\n"
                f"감독관을 최소 {3 * R}명 이상으로 입력해 주세요.")
            return
        try:
            res = solve(R, N, R5)
        except ValueError as e:
            messagebox.showerror("5교시 배정 불가", str(e))
            return
        if res is None:
            messagebox.showerror("배정 실패", "조건을 만족하는 배정을 찾지 못했습니다. 다시 시도해 주세요.")
            return
        self.result = res
        self.render()

    def render(self):
        res = self.result
        R, N, layout, stats = res["R"], res["N"], res["layout"], res["stats"]

        for t in (self.tree1, self.tree2):
            t.delete(*t.get_children())

        R5 = res["R5"]
        self.tree1["displaycolumns"] = self.cols1 if R5 else self.cols1[:-1]
        self.tree2["displaycolumns"] = [c for c in self.cols2 if c != "p5" or R5]

        for r in range(max(R, R5)):
            vals = [f"{r + 1}실"]
            for p in range(MAX_PERIODS):
                if p < len(layout) and r < len(layout[p]):
                    j, bs = layout[p][r]
                    vals.append(f"정 {j + 1}  ·  부 " + ", ".join(str(b + 1) for b in bs))
                else:
                    vals.append("-")
            self.tree1.insert("", "end", values=vals, tags=("odd" if r % 2 else "even",))

        for i, s in enumerate(stats):
            tags = ("warn",) if s["consec3"] else (("odd",) if i % 2 else ("even",))
            note = note_of(s)
            self.tree2.insert("", "end", iid=str(i + 1),
                              values=(f"{i + 1}번", s["jung"], s["bu"], s["jung"] + s["bu"],
                                      *s["cells"], note), tags=tags)

        jc = [s["jung"] for s in stats]
        cons_n = sum(1 for s in stats if s["consec3"])
        cons_j = [s["jung"] for s in stats if s["consec3"]]
        self.tile_vals["slots"].config(text=f"{sum(s['jung'] + s['bu'] for s in stats)}석")
        self.tile_vals["jung"].config(text=f"{min(jc)} ~ {max(jc)}회")
        self.tile_vals["cons"].config(
            text=f"{cons_n}명" + (f"  (정 {max(cons_j)}회 이하)" if cons_n else ""))
        self.tile_vals["p4"].config(text=f"{len(res['attend'][3])} / {N}명")
        self.tile_vals["p5"].config(
            text=f"{res['p5_overlap']} / {R5}명" if R5 else "미실시")

        msg = f"시험실 {R}개 · 감독관 {N}명 배정 완료" + (f" · 5교시 {R5}개 시험실" if R5 else "")
        if N != 3 * R:
            msg += "  |  감독관 수가 시험실×3과 달라 '전원 4교시 참여 · 1인 3회'는 불가 → 횟수를 최대한 균등 분배"
        self.status.config(text=msg)

    def find_staff(self):
        iid = self.find_var.get().strip()
        if iid in self.tree2.get_children():
            self.tree2.selection_set(iid)
            self.tree2.focus(iid)
            self.tree2.see(iid)
        else:
            messagebox.showinfo("찾기", "해당 번호의 감독관이 없습니다.")

    def export_excel(self):
        if not self.result:
            return
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            messagebox.showerror(
                "openpyxl 필요",
                "엑셀 저장에는 openpyxl 라이브러리가 필요합니다.\n"
                "터미널(명령 프롬프트)에서  pip install openpyxl  실행 후 다시 시도해 주세요.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel 파일", "*.xlsx")],
            initialfile="감독관_배정표.xlsx")
        if not path:
            return
        try:
            write_excel(path, self.result)
        except PermissionError:
            messagebox.showerror("저장 실패", "같은 이름의 파일이 엑셀에서 열려 있습니다.\n닫고 다시 시도해 주세요.")
            return
        messagebox.showinfo("저장 완료", f"엑셀 파일이 저장되었습니다.\n{path}")

    def export_csv(self):
        if not self.result:
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".csv", filetypes=[("CSV 파일", "*.csv")],
            initialfile="감독관_배정표.csv")
        if not path:
            return
        res = self.result
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["[교시별 배정표]"])
            w.writerow(["시험실", "1교시(정/부)", "2교시(정/부)", "3교시(정/부)",
                        "4교시(정/부/부)", "5교시(정/부)"])
            lay = res["layout"]
            for r in range(max(res["R"], res["R5"])):
                row = [f"{r + 1}실"]
                for p in range(MAX_PERIODS):
                    if p < len(lay) and r < len(lay[p]):
                        j, bs = lay[p][r]
                        row.append(f"정 {j + 1} / 부 " + ", ".join(str(b + 1) for b in bs))
                    else:
                        row.append("-")
                w.writerow(row)
            w.writerow([])
            w.writerow(["[감독관별 통계]"])
            w.writerow(["감독관", "정감독", "부감독", "합계", "1교시", "2교시", "3교시",
                        "4교시", "5교시", "비고"])
            for i, s in enumerate(res["stats"]):
                w.writerow([f"{i + 1}번", s["jung"], s["bu"], s["jung"] + s["bu"],
                            *s["cells"], note_of(s)])
        messagebox.showinfo("저장 완료", f"저장되었습니다.\n{path}")


if __name__ == "__main__":
    App().mainloop()
